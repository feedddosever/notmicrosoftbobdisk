"""Generate the synthetic customer workbook and the T00 stale-cache probe book.

Example Mutual Insurance Co. (FICTIONAL) - HO-3 rating workbook. Every name, rate and rule
is invented for a software demonstration. Policies carry synthetic IDs only.

usage: python -m tools.gen_workbook [--n 40] [--seed 2026] [--out workbook/example_mutual_ho3_rater.xlsx]

Sheets (layout fixed in sheetshift.json and docs/CONTRACT.md):
  Policies   : row 1 = the 14 input names, rows 2..n+1 = one policy per row
  RateTables : lookup tables + 13 defined names (A1 carries the FICTIONAL banner)
  Calc       : 43 formula columns A..AQ, row 1 = canonical output names, row-aligned with Policies
  Summary    : A1 banner, B2..B6 whole-book aggregates
  About      : the FICTIONAL banner and provenance (last sheet, so no row layout shifts)
Policies!A1 and Calc!A1 hold headers, so the banner is attached there as a cell comment.

Seeded spreadsheet anomalies (deliberate, disclosed; decided by a person, never by the tools):
  A1 Calc!O (all rows): RateTables!$A$31:$B$35 stops one row short of DedBands (A31:B36),
     so the $10,000 band (0.66) is never reached.
  A2 Calc!AM17: the number 48.17 typed over the tax formula.
  A3 Calc!X31: =W31, dropping the 25% credit cap.

Outputs are byte-deterministic for a given (n, seed): zip entry times and document
properties are pinned (see normalize_xlsx).

Author: Claude Code (AI agent) — scaffold; see ATTRIBUTION.md
"""
import argparse
import datetime as dt
import hashlib
import io
import os
import random
import re
import zipfile

import openpyxl
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName

CARRIER = "Example Mutual Insurance Co. (FICTIONAL)"
BANNER = "FICTIONAL — all names, rates and rules invented for a software demonstration"
FIXED_TIME = dt.datetime(2026, 1, 1, 0, 0, 0)

# ---------------------------------------------------------------- rate tables (fictional)
BASE_RATES = [  # zone, AOP rate /1000, hurricane rate /1000, premium tax rate
    ("T01", 3.125, 0.450, 0.0175), ("T02", 3.375, 0.875, 0.0175),
    ("T03", 3.640, 1.250, 0.0200), ("T04", 2.980, 2.125, 0.0200),
    ("T05", 4.115, 3.375, 0.0225), ("T06", 3.505, 4.625, 0.0225),
    ("T07", 4.250, 5.875, 0.0250), ("T08", 2.875, 6.500, 0.0250),
]
CONSTRUCTION = [("Frame", 1.00, 1.00), ("Masonry", 0.88, 0.80),
                ("MasonryVeneer", 0.94, 0.90), ("FireResistive", 0.80, 0.70)]
PROTECTION = [(1, .85, .82), (2, .87, .84), (3, .90, .86), (4, .93, .89), (5, .97, .93),
              (6, 1.00, .96), (7, 1.06, 1.01), (8, 1.15, 1.09), (9, 1.40, 1.30), (10, 1.75, 1.60)]
AGE_BANDS = [(0, .90), (5, .95), (10, 1.00), (20, 1.08), (30, 1.15), (50, 1.25), (75, 1.35)]
ROOF_BANDS = [(0, .95), (6, 1.00), (11, 1.10), (16, 1.25), (21, 1.45)]
AOI_BANDS = [(0, .80), (150000, .90), (250000, 1.00), (400000, 1.12), (600000, 1.25),
             (1000000, 1.40), (2000000, 1.55)]
DED_BANDS = [(0, 1.10), (500, 1.00), (1000, .92), (2500, .82), (5000, .74), (10000, .66)]
CLAIMS = [(0, 1.00), (1, 1.10), (2, 1.25), (3, 1.50)]
HURR_DED = [(0.02, 1.00), (0.05, .85), (0.10, .72)]
WIND_MIT = [("None", 1.00), ("Basic", .85), ("Fortified", .65)]
FEES = [(6, 15), (12, 25)]
CREDIT_NAMES = ["Alarm", "ClaimsFree", "NewHome", "MitBasic", "MitFortified"]
CREDIT_PCTS = [0.05, 0.10, 0.08, 0.04, 0.10]
SCALARS = {"CreditCap": 0.25, "MinPremium": 350, "AssessRate": 0.013, "HurrCapPct": 0.006}

INPUT_COLS = ["policy_id", "zone", "construction", "protection_class", "year_built", "roof_age",
              "coverage_a", "deductible", "hurr_ded_pct", "alarm", "wind_mit", "claims_3yr",
              "effective_date", "term_months"]  # Policies columns A..N, in this order

# ---------------------------------------------------------------- Calc formulas (row template)
# (column letter, canonical output name, formula template with {r} = sheet row)
CALC = [
    ("A", "policy_id", "=Policies!A{r}"),
    ("B", "home_age", '=IFERROR(DATEDIF(DATE(Policies!E{r},1,1),Policies!M{r},"y"),0)'),
    ("C", "roof_age_used", "=Policies!F{r}*1"),
    ("D", "aoi_units", "=Policies!G{r}/1000"),
    ("E", "base_rate", "=VLOOKUP(Policies!B{r},BaseRates,2,FALSE)"),
    ("F", "hurr_rate", "=VLOOKUP(Policies!B{r},BaseRates,3,FALSE)"),
    ("G", "tax_rate", "=VLOOKUP(Policies!B{r},BaseRates,4,FALSE)"),
    ("H", "base_premium", "=ROUND(E{r}*D{r},2)"),
    ("I", "constr_factor", "=INDEX(RateTables!$B$13:$B$16,MATCH(Policies!C{r},RateTables!$A$13:$A$16,0))"),
    ("J", "constr_hurr_factor", "=INDEX(RateTables!$C$13:$C$16,MATCH(Policies!C{r},RateTables!$A$13:$A$16,0))"),
    ("K", "pc_factor", '=INDEX(RateTables!$B$19:$C$28,MATCH(Policies!D{r},RateTables!$A$19:$A$28,0),IF(Policies!C{r}="Frame",1,2))'),
    ("L", "age_factor", "=VLOOKUP(B{r},AgeBands,2,TRUE)"),
    ("M", "roof_factor", "=VLOOKUP(C{r},RoofBands,2,TRUE)"),
    ("N", "aoi_factor", "=VLOOKUP(Policies!G{r},AOIBands,2,TRUE)"),
    ("O", "ded_factor", "=VLOOKUP(Policies!H{r},RateTables!$A$31:$B$35,2,TRUE)"),  # A1 anomaly
    ("P", "claims_factor", "=VLOOKUP(MIN(Policies!L{r},3),ClaimsTable,2,FALSE)"),
    ("Q", "aop_premium", "=ROUND(H{r}*I{r}*K{r}*L{r}*M{r}*N{r}*O{r}*P{r},2)"),
    ("R", "alarm_flag", '=IF(Policies!J{r}="Y",1,0)'),
    ("S", "claims_free_flag", "=IF(AND(Policies!L{r}=0,B{r}>=3),1,0)"),
    ("T", "new_home_flag", "=IF(B{r}<=5,1,0)"),
    ("U", "mit_basic_flag", '=IF(Policies!K{r}="Basic",1,0)'),
    ("V", "mit_fort_flag", '=IF(Policies!K{r}="Fortified",1,0)'),
    ("W", "credit_raw", "=SUMPRODUCT(R{r}:V{r},CreditPcts)"),
    ("X", "credit_pct", "=MIN(W{r},CreditCap)"),
    ("Y", "aop_net", "=ROUND(Q{r}*(1-X{r}),2)"),
    ("Z", "hurr_pct_used", '=IF(Policies!I{r}="",0.02,Policies!I{r})'),
    ("AA", "hurr_ded_factor", "=IFERROR(VLOOKUP(Z{r},HurrDedTable,2,FALSE),1)"),
    ("AB", "wind_mit_factor", "=IFERROR(INDEX(RateTables!$L$31:$L$33,MATCH(Policies!K{r},RateTables!$K$31:$K$33,0)),1)"),
    ("AC", "hurr_premium", "=ROUND(F{r}*D{r}*J{r}*AA{r}*AB{r},2)"),
    ("AD", "hurr_capped", "=MIN(AC{r},Policies!G{r}*HurrCapPct)"),
    ("AE", "subtotal", "=Y{r}+AD{r}"),
    ("AF", "exp_date", "=EDATE(Policies!M{r},Policies!N{r})"),
    ("AG", "term_factor", "=ROUND(YEARFRAC(Policies!M{r},AF{r},3),4)"),
    ("AH", "term_premium", "=ROUND(AE{r}*AG{r},2)"),
    ("AI", "min_applied", "=MAX(AH{r},MinPremium)"),
    ("AJ", "premium_rounded", "=ROUNDUP(AI{r},0)"),
    ("AK", "policy_fee", "=VLOOKUP(Policies!N{r},FeeTable,2,FALSE)"),
    ("AL", "assessment", "=ROUND(AJ{r}*AssessRate,2)"),
    ("AM", "tax", "=ROUND((AJ{r}+AK{r})*G{r},2)"),
    ("AN", "total_due", "=AJ{r}+AK{r}+AL{r}+AM{r}"),
    ("AO", "refer_flag", '=IF(OR(Policies!G{r}>1000000,C{r}>20,Policies!L{r}>=3),"REFER","OK")'),
    ("AP", "rate_per_1000", "=ROUND(AH{r}/D{r},3)"),
    ("AQ", "rate_class", '=Policies!B{r}&"-"&LEFT(Policies!C{r},1)&TEXT(Policies!D{r},"00")'),
]
CALC_NAMES = [c[1] for c in CALC]
ANOMALY_HARDCODE = ("AM", 17, 48.17)        # A2: value typed over the tax formula
ANOMALY_INCONSISTENT = ("X", 31, "=W{r}")   # A3: credit cap dropped in one row

# Summary!B2..B6 (row 1 = banner); {last} = last data row
SUMMARY = [
    ("Policies", "=COUNTA(Calc!A2:A{last})"),
    ("Written premium", "=SUM(Calc!AJ2:AJ{last})"),
    ("Total due", "=SUM(Calc!AN2:AN{last})"),
    ("Avg rate/1000", "=AVERAGE(Calc!AP2:AP{last})"),
    ("Hurricane share", "=SUMPRODUCT(Calc!AD2:AD{last})/SUM(Calc!AE2:AE{last})"),
]

NAMES = {
    "BaseRates": "RateTables!$A$3:$D$10",
    "AgeBands": "RateTables!$E$12:$F$18",
    "RoofBands": "RateTables!$H$12:$I$16",
    "AOIBands": "RateTables!$K$12:$L$18",
    "DedBands": "RateTables!$A$31:$B$36",   # the table the formula SHOULD use (see A1)
    "ClaimsTable": "RateTables!$E$31:$F$34",
    "HurrDedTable": "RateTables!$H$31:$I$33",
    "FeeTable": "RateTables!$N$31:$O$32",
    "CreditPcts": "RateTables!$B$40:$F$40",
    "CreditCap": "RateTables!$B$42",
    "MinPremium": "RateTables!$B$43",
    "AssessRate": "RateTables!$B$44",
    "HurrCapPct": "RateTables!$B$45",
}


def _write_rate_tables(ws):
    ws["A1"] = f"{BANNER} | {CARRIER} HO-3 rate tables"
    for i, h in enumerate(["Zone", "AOP rate/1000", "Hurr rate/1000", "Tax rate"]):
        ws.cell(2, 1 + i, h)
    for k, row in enumerate(BASE_RATES):
        for j, v in enumerate(row):
            ws.cell(3 + k, 1 + j, v)
    ws["A12"], ws["B12"], ws["C12"] = "Construction", "AOP factor", "Hurr factor"
    for k, row in enumerate(CONSTRUCTION):
        for j, v in enumerate(row):
            ws.cell(13 + k, 1 + j, v)
    ws["A18"], ws["B18"], ws["C18"] = "PC", "Frame", "NonFrame"
    for k, row in enumerate(PROTECTION):
        for j, v in enumerate(row):
            ws.cell(19 + k, 1 + j, v)
    for col, title, table in ((5, "Home age >=", AGE_BANDS), (8, "Roof age >=", ROOF_BANDS),
                              (11, "Coverage A >=", AOI_BANDS)):
        ws.cell(11, col, title)
        for k, (a, b) in enumerate(table):
            ws.cell(12 + k, col, a)
            ws.cell(12 + k, col + 1, b)
    for col, title, table in ((1, "Deductible >=", DED_BANDS), (5, "Claims", CLAIMS),
                              (8, "Hurr ded %", HURR_DED), (11, "Wind mitigation", WIND_MIT),
                              (14, "Term months", FEES)):
        ws.cell(30, col, title)
        for k, (a, b) in enumerate(table):
            ws.cell(31 + k, col, a)
            ws.cell(31 + k, col + 1, b)
    ws["A39"] = "Credit"
    for j, (n, p) in enumerate(zip(CREDIT_NAMES, CREDIT_PCTS)):
        ws.cell(39, 2 + j, n)
        ws.cell(40, 2 + j, p)
    ws["A40"] = "Credit %"
    for k, (n, v) in enumerate(SCALARS.items()):
        ws.cell(42 + k, 1, n)
        ws.cell(42 + k, 2, v)


def _write_about(ws, n, seed):
    rows = [
        (BANNER, None),
        (None, None),
        ("Carrier", CARRIER),
        ("Product", "Homeowners HO-3 rating workbook (synthetic demonstration)"),
        ("Generated by", f"tools/gen_workbook.py --n {n} --seed {seed}"),
        ("Layout", f"Policies rows 2-{n + 1}; Calc formula columns A:AQ rows 2-{n + 1}; Summary B2:B6"),
        ("Disclosure", "Contains deliberately seeded spreadsheet anomalies for the demonstration "
                       "(see docs/CONTRACT.md)."),
        ("Personal data", "None. Policy IDs are synthetic; there are no names, addresses or contacts."),
    ]
    for r, (a, b) in enumerate(rows, 1):
        if a is not None:
            ws.cell(r, 1, a)
        if b is not None:
            ws.cell(r, 2, b)


def build(policies, path, anomalies=True, seed=2026):
    """Write the workbook for `policies` (list of dicts keyed by INPUT_COLS) to `path`."""
    wb = openpyxl.Workbook()
    wb.properties.creator = "SheetShift tools/gen_workbook.py"
    wb.properties.title = f"{CARRIER} HO-3 rater (synthetic)"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    wp = wb.active
    wp.title = "Policies"
    wp.append(INPUT_COLS)
    wp["A1"].comment = Comment(BANNER, "SheetShift")
    for p in policies:
        wp.append([p.get(c) for c in INPUT_COLS])
    n = len(policies)
    for r in range(2, n + 2):
        wp.cell(r, 13).number_format = "yyyy-mm-dd"
    _write_rate_tables(wb.create_sheet("RateTables"))
    for name, ref in NAMES.items():
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    wc = wb.create_sheet("Calc")
    wc.append(CALC_NAMES)
    wc["A1"].comment = Comment(BANNER, "SheetShift")
    for r in range(2, n + 2):
        for col, _, f in CALC:
            wc[f"{col}{r}"] = f.format(r=r)
        wc[f"AF{r}"].number_format = "yyyy-mm-dd"
    if anomalies:
        col, row, val = ANOMALY_HARDCODE
        if row <= n + 1:
            wc[f"{col}{row}"] = val
        col, row, f = ANOMALY_INCONSISTENT
        if row <= n + 1:
            wc[f"{col}{row}"] = f.format(r=row)
    ws = wb.create_sheet("Summary")
    ws["A1"] = BANNER
    for k, (label, f) in enumerate(SUMMARY):
        ws.cell(2 + k, 1, label)
        ws.cell(2 + k, 2, f.format(last=n + 1))
    _write_about(wb.create_sheet("About"), n, seed)
    wb.save(path)
    normalize_xlsx(path)


def normalize_xlsx(path):
    """Rewrite an xlsx in place with pinned zip entry times and document-property dates."""
    stamp = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ")
    with zipfile.ZipFile(path) as z:
        entries = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in entries:
            if name == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)",
                              lambda m: m.group(1) + stamp + m.group(3), text)
                data = text.encode("utf-8")
            info = zipfile.ZipInfo(name, date_time=FIXED_TIME.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            out.writestr(info, data)
    with open(path, "wb") as f:
        f.write(buf.getvalue())


def rate_tables_json():
    """The generator's own view of the tables (the mapping tool extracts them from the book)."""
    return {
        "base_rates": {z: {"aop": a, "hurr": h, "tax": t} for z, a, h, t in BASE_RATES},
        "construction": {c: {"aop": a, "hurr": h} for c, a, h in CONSTRUCTION},
        "protection": {str(p): {"frame": f, "nonframe": nf} for p, f, nf in PROTECTION},
        "age_bands": AGE_BANDS, "roof_bands": ROOF_BANDS, "aoi_bands": AOI_BANDS,
        "ded_bands": DED_BANDS, "claims": CLAIMS, "hurr_ded": HURR_DED,
        "wind_mit": WIND_MIT, "fees": FEES,
        "credits": dict(zip(CREDIT_NAMES, CREDIT_PCTS)), **SCALARS,
    }


# ---------------------------------------------------------------- policy generation
ZONES = [z[0] for z in BASE_RATES]


def random_policy(rng, i):
    """One synthetic policy; blanks (None) and case variants are deliberate (traps T3/T4)."""
    eff = dt.date(2026, 1, 1) + dt.timedelta(days=rng.randrange(0, 900))
    return {
        "policy_id": f"HO-{i:06d}",
        "zone": rng.choice(ZONES),
        "construction": rng.choice([c[0] for c in CONSTRUCTION]),
        "protection_class": rng.randint(1, 10),
        "year_built": rng.randint(1925, eff.year),
        "roof_age": rng.choice([None] + list(range(0, 30))),
        "coverage_a": rng.choice([rng.randrange(100, 1600) * 1000, rng.randrange(100000, 1600000)]),
        "deductible": rng.choice([500, 1000, 1000, 2500, 2500, 5000, 10000, 250, 1500, 7500, 25000]),
        "hurr_ded_pct": rng.choice([None, 0.02, 0.05, 0.10, 0.03]),
        "alarm": rng.choice(["Y", "N", None, "y", "Y"]),
        "wind_mit": rng.choice(["None", "Basic", "Fortified", None, "fortified"]),
        "claims_3yr": rng.choice([None, 0, 0, 0, 1, 1, 2, 3, 4]),
        "effective_date": eff,
        "term_months": rng.choice([12, 12, 12, 6]),
    }


# Row-specific overrides (sheet row -> fields) so the small customer book exercises the
# seeded anomalies and the error path. Row 31 exceeds the 25% credit cap (so A3 shows at
# runtime); row 17 is an ordinary valid policy under the typed-over tax (A2); row 40 uses an
# ineligible territory, so #N/A propagates to total_due (trap T8).
CURATED = {
    17: {"zone": "T03", "deductible": 1000, "term_months": 12},
    31: {"alarm": "Y", "wind_mit": "Fortified", "claims_3yr": 0, "year_built": 2023,
         "effective_date": dt.date(2027, 6, 1), "deductible": 1000},
    40: {"zone": "T09"},
}


def customer_policies(n, seed):
    """The n policies of the customer workbook, deterministic for (n, seed)."""
    rng = random.Random(seed)
    pols = [random_policy(rng, i + 1) for i in range(n)]
    for row, fields in CURATED.items():
        if row - 2 < n:
            pols[row - 2].update(fields)
    return pols


# ---------------------------------------------------------------- T00 probe book
def build_stale_cache_probe(path, cached_value=99):
    """A tiny book whose cached D2 value (99) differs from its formula result (=B2*C2 = 7.5).

    A reader that shows cached values reports 99; a forced recalculation reports 7.5.
    The app.xml generator stays "Microsoft Excel", so LibreOffice's default profile trusts
    (and keeps) the cache; the forced profile recalculates.
    """
    wb = openpyxl.Workbook()
    wb.properties.creator = "SheetShift tools/gen_workbook.py"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    ws = wb.active
    ws.title = "Sheet1"
    for c, h in enumerate(["item", "qty", "unit_price", "total"], 1):
        ws.cell(1, c, h)
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "widget", 3, 2.5, "=B2*C2"
    ws["A4"] = f"{BANNER} (probe workbook)"
    wb.save(path)
    with zipfile.ZipFile(path) as z:
        entries = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    patched = []
    for name, data in entries:
        if name == "xl/worksheets/sheet1.xml":
            text, k = re.subn(r'(<c r="D2"[^>]*>\s*<f>B2\*C2</f>)\s*(<v\s*/>|<v>\s*</v>)?',
                              lambda m: f"{m.group(1)}<v>{cached_value}</v>", data.decode("utf-8"))
            assert k == 1, "could not place the cached value in D2"
            data = text.encode("utf-8")
        elif name == "xl/workbook.xml":
            data = re.sub(rb'\s*fullCalcOnLoad="1"', b"", data)
        patched.append((name, data))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in patched:
            out.writestr(name, data)
    normalize_xlsx(path)


def sha256(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def write_sums(paths, out_path, root):
    lines = [f"{sha256(p)}  {os.path.relpath(p, root)}" for p in sorted(paths)]
    with open(out_path, "w", newline="\n") as f:
        f.write("\n".join(lines) + "\n")


def main(argv=None):
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--n", type=int, default=40)
    ap.add_argument("--seed", type=int, default=2026)
    ap.add_argument("--out", default=os.path.join(root, "workbook", "example_mutual_ho3_rater.xlsx"))
    ap.add_argument("--probe", default=os.path.join(root, "workbook", "probe", "stale_cache.xlsx"))
    ap.add_argument("--no-sums", action="store_true", help="skip workbook/SHA256SUMS")
    a = ap.parse_args(argv)
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    os.makedirs(os.path.dirname(a.probe), exist_ok=True)
    build(customer_policies(a.n, a.seed), a.out, seed=a.seed)
    build_stale_cache_probe(a.probe)
    if not a.no_sums:
        write_sums([a.out, a.probe], os.path.join(os.path.dirname(a.out), "SHA256SUMS"),
                   os.path.dirname(a.out))
    print(f"wrote {a.out} ({a.n} policies, {a.n * len(CALC)} Calc formula cells) and {a.probe}")


if __name__ == "__main__":
    main()
