"""Map the customer workbook into the committed, deterministic build/ outputs.

usage: python -m tools.dump_workbook [--workbook PATH] [--out build] [--recalc] [--expect-lints N]

Writes (layout in docs/CONTRACT.md):
  build/graph.json            43 Calc column rules + Summary cells, topo order, names, exceptions
  build/units.json            U1..U4: columns, outputs, inputs, upstream
  build/lints.json            spreadsheet-anomaly lints with ids, column rule and manual rule
  build/rate_tables.json      every lookup table and scalar, extracted from the workbook
  build/workbook_values.json  the workbook's own values (LibreOffice forced recalc), keyed by sha256
  build/sheets/Calc.md        the 43 column rules (< 15k chars)
  build/sheets/Calc_exceptions.md
  build/units/U1..U4.md       one brief per unit with 5 sample rows (< 5k chars each)

LibreOffice is needed only when build/workbook_values.json does not match the workbook's
sha256 (or with --recalc); otherwise the committed values are reused, so members without
LibreOffice can run the map. Flags volatile/unsupported constructs (NOW, TODAY, RAND,
OFFSET, INDIRECT, macros, external links, array formulas, iterative calculation).

Author: Claude Code (AI agent) — scaffold; see ATTRIBUTION.md
"""
import argparse
import datetime as dt
import hashlib
import json
import os
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

if __package__ in (None, ""):  # run as `python3 tools/dump_workbook.py`: make `tools` importable
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from tools import depgraph, units  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CALC_MD_LIMIT = 15000
LINT_ORDER = ["range_short_of_table", "hardcoded_value_in_formula_column", "inconsistent_formula"]
LIMITS = [
    "Macros and VBA are not supported; .xlsm files are never opened.",
    "Volatile functions (NOW, TODAY, RAND, OFFSET, INDIRECT) are listed as unsupported.",
    "External links, array and dynamic-array formulas, data tables and iterative calculation are out of scope.",
    "The oracle is LibreOffice 24.2 with forced recalculation; Excel parity is UNVERIFIED unless the Excel cross-check runs.",
    "Inputs are sampled, so equivalence is evidence, not proof.",
    "The workbook's anomalies were seeded deliberately for the demonstration.",
]


def sha256(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def write(path, text):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)


FLAT_LIST = re.compile(r"\[\s*\n\s*([^\[\]{}]*?)\s*\n\s*\]")


def write_json(path, obj):
    """Stable JSON (insertion-ordered keys); lists of scalars are kept on one line."""
    text = json.dumps(obj, indent=1, ensure_ascii=False)
    text = FLAT_LIST.sub(lambda m: "[" + re.sub(r",\s*\n\s*", ", ", m.group(1)) + "]", text)
    write(path, text + "\n")


# ---------------------------------------------------------------- workbook reading
def read_layout_checked(wb, layout):
    """Fail loudly if the workbook's headers do not match sheetshift.json."""
    s = layout["sheets"]
    hdr_in = [c.value for c in wb[s["inputs"]][layout["header_row"]]][:len(layout["inputs"])]
    hdr_out = [c.value for c in wb[s["calc"]][layout["header_row"]]][:len(layout["outputs"])]
    if hdr_in != layout["inputs"]:
        sys.exit(f"Policies headers {hdr_in} do not match sheetshift.json inputs")
    if hdr_out != layout["outputs"]:
        sys.exit(f"Calc headers {hdr_out} do not match sheetshift.json outputs")


def read_inputs(wb, layout):
    """{sheet_row: {input_name: value}} with dates as datetime.date and blanks as None."""
    ws = wb[layout["sheets"]["inputs"]]
    out = {}
    for r in range(layout["first_data_row"], layout["last_data_row"] + 1):
        rec = {}
        for name, col in zip(layout["inputs"], layout["input_columns"]):
            v = ws[f"{col}{r}"].value
            rec[name] = v.date() if isinstance(v, dt.datetime) else v
        out[r] = rec
    return out


def encode(v):
    """Oracle value -> JSON: numbers, text, null, {"date": iso}, {"error": code}."""
    if isinstance(v, dt.date):
        return {"date": v.isoformat()}
    if units.is_error(v):
        return {"error": v}
    return v


def decode(v):
    if isinstance(v, dict):
        return dt.date.fromisoformat(v["date"]) if "date" in v else v["error"]
    return v


def workbook_values(path, sha, layout, cache_path, force):
    """The workbook's own Calc values: reuse the committed cache or recalculate once."""
    if not force and os.path.exists(cache_path):
        with open(cache_path, encoding="utf-8") as f:
            cached = json.load(f)
        if cached.get("workbook_sha256") == sha:
            return cached, {r["row"]: {k: decode(v) for k, v in r["values"].items()}
                            for r in cached["rows"]}
    from tools import lo_recalc
    if not lo_recalc.soffice():
        sys.exit("build/workbook_values.json is missing or stale and LibreOffice is not installed")
    header, recs = lo_recalc.recalc_sheet(path, layout["sheets"]["calc"], tuple(layout["date_outputs"]))
    first = layout["first_data_row"]
    rows = {first + i: rec for i, rec in enumerate(recs[:layout["n_rows"]])}
    cached = {
        "_about": "Values the workbook computes for itself (the oracle for the unit samples). FICTIONAL data.",
        "workbook": layout["workbook"], "workbook_sha256": sha,
        "engine": lo_recalc.version(), "recalc": "forced (OOXMLRecalcMode=0)",
        "sheet": layout["sheets"]["calc"], "encoding": "number | text | null | {date} | {error}",
        "rows": [{"row": r, "values": {k: encode(rows[r].get(k)) for k in layout["outputs"]}}
                 for r in sorted(rows)],
    }
    return cached, rows


# ---------------------------------------------------------------- graph composition
def compose_graph(g, layout, sha):
    s = layout["sheets"]
    first = layout["first_data_row"]
    in_by_col = dict(zip(layout["input_columns"], layout["inputs"]))
    out_by_col = dict(zip(layout["output_columns"], layout["outputs"]))
    rules = layout.get("manual_rules", {})

    def split_prec(node):
        calc, inp = [], []
        for sh, col in node["precedent_columns"]:
            if sh == s["calc"] and col in out_by_col and col != node["col"]:
                calc.append(col)
            elif sh == s["inputs"] and col in in_by_col:
                inp.append(in_by_col[col])
        return calc, [i for i in layout["inputs"] if i in inp]

    columns, cells = [], []
    for n in sorted(g["nodes"], key=lambda n: n["topo_index"]):
        calc, inp = split_prec(n)
        if n["sheet"] == s["calc"] and n["kind"] == "column":
            colidx = column_index_from_string(n["col"])
            columns.append({
                "col": n["col"], "cell": n["key"], "output_name": out_by_col[n["col"]],
                "rows": n["rows"], "n_cells": n["n_cells"],
                "template_r1c1": n["template_r1c1"],
                "template_a1": depgraph.a1(n["template_r1c1"], first, colidx),
                "precedents": [out_by_col[c] for c in calc], "precedent_cols": calc,
                "inputs": inp, "tables": n["tables"], "named_ranges": n["names"],
                "functions": n["functions"], "exceptions": n["exceptions"],
                "manual_rule": rules.get(out_by_col[n["col"]]), "topo_index": n["topo_index"],
            })
        else:
            row = int(re.search(r"(\d+)$", n["key"]).group(1))
            cells.append({
                "cell": n["key"], "sheet": n["sheet"], "label": None, "formula": n["formula_a1"],
                "template_r1c1": n["template_r1c1"], "row": row,
                "precedents": [out_by_col[c] for c in calc], "precedent_cols": calc, "inputs": inp,
                "named_ranges": n["names"], "functions": n["functions"], "topo_index": n["topo_index"],
            })
    columns.sort(key=lambda c: column_index_from_string(c["col"]))
    return columns, cells


def compose_lints(g, layout, columns):
    """Lints with stable ids A1.. (ordered by type, then column, then row)."""
    s = layout["sheets"]
    by_key = {c["cell"]: c for c in columns}
    first, last = layout["first_data_row"], layout["last_data_row"]
    out = []
    for l in g["lints"]:
        c = by_key.get(l["node"])
        if c is None:
            continue
        col = c["col"]
        base = {"id": None, "type": l["type"], "sheet": s["calc"], "column": col,
                "output_name": c["output_name"]}
        if l["type"] == "range_short_of_table":
            rows = l["rows"]
            base.update(cell=c["cell"], cells=f"{c['cell']}{min(rows)}:{col}{max(rows)}",
                        row=None, formula=c["template_a1"], column_rule=c["template_a1"],
                        range=l["range"], table_block=l["table_block"],
                        rows_missing=l["rows_missing"], missing_keys=l["missing_keys"],
                        named_ranges_covering_block=l["named_ranges_covering_block"])
        else:
            r = l["row"]
            rule = depgraph.a1(c["template_r1c1"], r, column_index_from_string(col))
            formula = l.get("formula") if l["type"] == "inconsistent_formula" else str(l["value"])
            base.update(cell=f"{c['cell']}{r}", cells=f"{c['cell']}{r}", row=r, formula=formula,
                        column_rule=rule)
            if "value" in l:
                base["value"] = l["value"]
        base["manual_rule"] = c["manual_rule"]
        base["options"] = layout.get("anomaly_options", {}).get(l["type"], ["adopt-manual", "escalate"])
        base["status"] = "open"
        assert first <= (base["row"] or first) <= last
        out.append(base)
    out.sort(key=lambda b: (LINT_ORDER.index(b["type"]) if b["type"] in LINT_ORDER else 99,
                            column_index_from_string(b["column"]), b["row"] or 0))
    for i, b in enumerate(out, 1):
        b["id"] = f"A{i}"
    return out


def affected_rows(lints, inputs, values, layout):
    """Sheet rows whose values a lint touches (left out of the unit samples)."""
    rows = set()
    in_by_col = dict(zip(layout["input_columns"], layout["inputs"]))
    out_by_col = dict(zip(layout["output_columns"], layout["outputs"]))
    for l in lints:
        if l["row"]:
            rows.add(l["row"])
            continue
        m = re.match(r"=VLOOKUP\((?:(\w+)!)?\$?([A-Z]+)\$?\d+,", l["formula"] or "")
        if not (m and l["missing_keys"]):
            rows.update(range(layout["first_data_row"], layout["last_data_row"] + 1))
            continue
        sheet, col = m.group(1) or layout["sheets"]["calc"], m.group(2)
        lo = min(k for k in l["missing_keys"] if isinstance(k, (int, float)))
        for r in inputs:
            v = inputs[r].get(in_by_col.get(col)) if sheet == layout["sheets"]["inputs"] \
                else values[r].get(out_by_col.get(col))
            if isinstance(v, (int, float)) and v >= lo:
                rows.add(r)
    return rows


def extract_tables(wb, g, names):
    """Every defined name and every literal lookup range, with the header row above it."""
    def block(sheet, c1, r1, c2, r2):
        ws = wb[sheet]
        hdr = [ws.cell(r1 - 1, c).value for c in range(c1, c2 + 1)] if r1 > 1 else []
        rows = [[ws.cell(r, c).value for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]
        return hdr, rows

    tables, scalars = {}, {}
    spans = []
    for nm in sorted(names):
        sheet, ref = names[nm].split("!")
        c1, r1, c2, r2 = range_boundaries(ref.replace("$", ""))
        spans.append((sheet, c1, r1, c2, r2))
        if (c1, r1) == (c2, r2):
            scalars[nm] = wb[sheet].cell(r1, c1).value
            continue
        hdr, rows = block(sheet, c1, r1, c2, r2)
        tables[nm] = {"ref": names[nm], "header": hdr, "rows": rows, "literal_refs": []}
    literal = {}
    for n in g["nodes"]:
        for t in n["tables"]:
            if t["name"]:
                continue
            sheet, ref = t["ref"].split("!")
            c1, r1, c2, r2 = range_boundaries(ref)
            inside = [nm for nm in tables if (sheet, ) == (names[nm].split("!")[0], ) and
                      _within((c1, r1, c2, r2), range_boundaries(names[nm].split("!")[1].replace("$", "")))]
            if inside:
                ref_note = f"{t['operand']} ({n['key']})"
                if ref_note not in tables[inside[0]]["literal_refs"]:
                    tables[inside[0]]["literal_refs"].append(ref_note)
                continue
            literal.setdefault((sheet, r1, r2), set()).update(range(c1, c2 + 1))
    for (sheet, r1, r2), cols in sorted(literal.items()):
        cols = sorted(cols)
        runs, start = [], cols[0]
        for a, b in zip(cols, cols[1:] + [None]):
            if b != a + 1:
                runs.append((start, a))
                start = b
        for c1, c2 in runs:
            key = f"{sheet}!${get_column_letter(c1)}${r1}:${get_column_letter(c2)}${r2}"
            hdr, rows = block(sheet, c1, r1, c2, r2)
            tables[key] = {"ref": key, "header": hdr, "rows": rows, "literal_refs": []}
    for t in tables.values():
        if not t["literal_refs"]:
            del t["literal_refs"]
    return tables, scalars


def _within(inner, outer):
    return outer[0] <= inner[0] and outer[1] <= inner[1] and inner[2] <= outer[2] and inner[3] <= outer[3]


# ---------------------------------------------------------------- markdown
def calc_md(columns, layout, names, unit_map, sha):
    in_pairs = " · ".join(f"{c} {n}" for c, n in zip(layout["input_columns"], layout["inputs"]))
    first, last = layout["first_data_row"], layout["last_data_row"]
    L = [f"# Calc: {len(columns)} column rules", "",
         f"{layout['carrier']}: synthetic workbook `{layout['workbook']}` (sha256 {sha[:12]}). "
         "Generated by tools/dump_workbook.py; do not edit.", "",
         f"Every data row {first}..{last} of Calc uses its column's rule, shown here for row {first}, "
         f"except the cells in Calc_exceptions.md. Calc row r reads Policies row r.", "",
         f"Inputs (Policies): {in_pairs}", "",
         "Named ranges: " + " · ".join(f"{k}={v}" for k, v in names.items()), "",
         "Units: " + " · ".join(f"{u} {v['range'].split('!')[1]}" for u, v in unit_map.items()), "",
         "Summary!B2..B6 are whole-book aggregates (build/graph.json `cells`), not part of quote().", "",
         "| Col | output_name | Rule (row %d) | Reads | Ref |" % first, "|---|---|---|---|---|"]
    for c in columns:
        reads = [f"p.{i}" for i in c["inputs"]] + [f"c.{p}" for p in c["precedents"]]
        L.append(f"| {c['col']} | {c['output_name']} | `{c['template_a1']}` | "
                 f"{', '.join(reads) or '-'} | {c['manual_rule'] or '-'} |")
    return "\n".join(L) + "\n"


def exceptions_md(lints, unsupported, layout, sha):
    L = ["# Calc: exceptions and lints", "",
         f"{layout['carrier']}: synthetic workbook (sha256 {sha[:12]}). Generated by "
         "tools/dump_workbook.py; do not edit.", "",
         "The anomalies in this demonstration workbook were seeded deliberately. The tools only "
         "detect them; a person decides each one (decisions/decisions.jsonl), citing the manual.", ""]
    for l in lints:
        L.append(f"## {l['id']}: {l['type']} at {l['cells']} ({l['output_name']})")
        L.append("")
        if l["type"] == "range_short_of_table":
            L.append(f"- Formula (row {layout['first_data_row']}): `{l['formula']}`")
            L.append(f"- Range {l['range']} stops {l['rows_missing']} row(s) short of the table "
                     f"block {l['table_block']} (named range: {', '.join(l['named_ranges_covering_block']) or 'none'}); "
                     f"unreachable key(s): {', '.join(str(k) for k in l['missing_keys'])}")
        else:
            L.append(f"- Cell content: `{l['formula']}`")
            L.append(f"- Column rule for that row: `{l['column_rule']}`")
        L.append(f"- Manual rule: {l['manual_rule']} · allowed options: {', '.join(l['options'])} · status: {l['status']}")
        L.append("")
    L.append("## Unsupported constructs")
    L.append("")
    if unsupported:
        L += [f"- {u['kind']}: {u['detail']}" for u in unsupported]
    else:
        L.append("- None found (no volatile functions, macros, external links, array formulas "
                 "or iterative calculation).")
    return "\n".join(L) + "\n"


# ---------------------------------------------------------------- main
def main(argv=None):
    ap = argparse.ArgumentParser(description="Map the workbook into build/")
    ap.add_argument("--layout", default=os.path.join(ROOT, "sheetshift.json"))
    ap.add_argument("--workbook", default=None)
    ap.add_argument("--out", default=os.path.join(ROOT, "build"))
    ap.add_argument("--recalc", action="store_true", help="recalculate with LibreOffice even if cached")
    ap.add_argument("--expect-lints", type=int, default=None)
    a = ap.parse_args(argv)
    with open(a.layout, encoding="utf-8") as f:
        layout = json.load(f)
    path = a.workbook or os.path.join(ROOT, layout["workbook"])
    if path.lower().endswith((".xlsm", ".xlsb", ".xls")):
        sys.exit("only .xlsx workbooks are supported (macros are never opened)")
    sha = sha256(path)
    wb = openpyxl.load_workbook(path)
    read_layout_checked(wb, layout)
    g = depgraph.analyze(path)
    names = dict(sorted(g["names"].items()))
    columns, cells = compose_graph(g, layout, sha)
    summary_ws = wb[layout["sheets"]["summary"]]
    for c in cells:
        if c["sheet"] == layout["sheets"]["summary"]:
            c["label"] = summary_ws.cell(c["row"], 1).value
    lints = compose_lints(g, layout, columns)
    order = [c["col"] for c in sorted(columns, key=lambda c: c["topo_index"])]
    parts = units.split(order, layout["units"]["count"], layout["units"].get("target_starts"))
    unit_map = units.describe(parts, columns, layout["inputs"])
    for uid, u in unit_map.items():
        u["lints"] = [l["id"] for l in lints if l["column"] in u["columns"]]
    owner = {col: uid for uid, u in unit_map.items() for col in u["columns"]}
    for c in columns:
        c["unit"] = owner[c["col"]]

    cache_path = os.path.join(a.out, "workbook_values.json")
    cached, values = workbook_values(path, sha, layout, cache_path, a.recalc)
    inputs = read_inputs(wb, layout)
    rows = {r: {"p": inputs[r], "c": values[r]} for r in inputs}
    excluded = affected_rows(lints, inputs, values, layout)

    graph = {
        "_about": f"{layout['carrier']}: dependency graph of the synthetic workbook. Generated; do not edit.",
        "workbook": layout["workbook"], "workbook_sha256": sha, "sheets": g["sheets"],
        "layout": {k: layout[k] for k in ("header_row", "first_data_row", "last_data_row", "n_rows")},
        "inputs": [{"name": n, "col": c} for n, c in zip(layout["inputs"], layout["input_columns"])],
        "columns": columns, "cells": cells,
        "topo_order": [c["output_name"] for c in sorted(columns, key=lambda c: c["topo_index"])],
        "rule_order": [x["cell"] for x in sorted(columns + cells, key=lambda x: x["topo_index"])],
        "named_ranges": names,
        "exceptions": [{"id": l["id"], "type": l["type"], "cell": l["cells"]} for l in lints],
        "function_inventory": depgraph.function_inventory(g["nodes"]),
        "unsupported": g["unsupported"], "limits": LIMITS,
        "stats": {"formula_cells": sum(n["n_cells"] for n in g["nodes"]),
                  "rules": len(columns) + len(cells), "calc_columns": len(columns),
                  "other_cells": len(cells)},
    }
    tables, scalars = extract_tables(wb, g, g["names"])
    rate_tables = {"_about": f"{layout['carrier']}: invented rates extracted from {layout['workbook']}. "
                             "Generated by tools/dump_workbook.py; do not edit.",
                   "workbook_sha256": sha, "tables": tables, "scalars": scalars}

    out = a.out
    write_json(os.path.join(out, "graph.json"), graph)
    write_json(os.path.join(out, "units.json"), unit_map)
    write_json(os.path.join(out, "lints.json"), {"_about": "Spreadsheet-anomaly lints; a person decides each.",
                                                 "workbook_sha256": sha, "count": len(lints), "lints": lints})
    write_json(os.path.join(out, "rate_tables.json"), rate_tables)
    write_json(cache_path, cached)
    errors = []
    text = calc_md(columns, layout, names, unit_map, sha)
    if len(text) >= CALC_MD_LIMIT:
        errors.append(f"Calc.md is {len(text)} chars (limit {CALC_MD_LIMIT})")
    write(os.path.join(out, "sheets", "Calc.md"), text)
    write(os.path.join(out, "sheets", "Calc_exceptions.md"), exceptions_md(lints, g["unsupported"], layout, sha))
    sizes = {"Calc.md": len(text)}
    for uid, u in unit_map.items():
        samples = units.choose_samples(u, rows, excluded)
        md = units.render_md(uid, u, columns, layout, names, lints, rows, samples, cached["engine"])
        if not units.check_size(md):
            errors.append(f"{uid}.md is {len(md)} chars (limit {units.UNIT_MD_LIMIT})")
        write(os.path.join(out, "units", f"{uid}.md"), md)
        sizes[f"{uid}.md"] = len(md)
    print(f"map: {len(columns)} columns, {len(cells)} other cells, {len(lints)} lints "
          f"({', '.join(l['id'] + ' ' + l['cells'] for l in lints)}), unsupported={len(g['unsupported'])}")
    print("sizes: " + ", ".join(f"{k}={v}" for k, v in sizes.items()))
    if a.expect_lints is not None and len(lints) != a.expect_lints:
        errors.append(f"expected {a.expect_lints} lints, found {len(lints)}")
    if errors:
        sys.exit("map FAILED: " + "; ".join(errors))


if __name__ == "__main__":
    main()
