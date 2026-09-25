"""Apply signed adopt-manual decisions to a copy of the workbook, then expand it.

usage: python -m harness.patch_workbook [--seed 2026]
       python -m harness.oracle_lo --seed 2026 --patched      (then recalculate it)

Reads decisions/decisions.jsonl (written by people with tools/decide.py). For every
adopt-manual decision it applies the deterministic repair for the lint's type, after
checking that the decision cites the manual rule the repair implements:
  range_short_of_table              (R-205)  every cell of the column: the short range is
                                             replaced by the named range covering the table
                                             block (DedBands), or by the block itself
  hardcoded_value_in_formula_column (R-510)  the typed value becomes the column formula again
  inconsistent_formula              (R-310)  the odd cell gets the column formula again
keep-workbook and escalate decisions change nothing. The patched customer book goes to
build/cache/harness/customer_patched.xlsx and is expanded with the golden inputs to
build/cache/harness/expanded_<seed>_patched.xlsx; a manifest records what changed.

Needs openpyxl: Claude Code sandbox / CI only. Never invents decisions: with no
decisions file it does nothing.

Author: Claude Code (AI agent) — scaffold; see ATTRIBUTION.md
"""
import argparse
import json
import os
import sys

from harness import common as C
from harness import expand as E

REPAIR_RULES = {"range_short_of_table": "R-205", "hardcoded_value_in_formula_column": "R-510",
                "inconsistent_formula": "R-310"}


def _cells(ws, col, first, last):
    return [ws["%s%d" % (col, r)] for r in range(first, last + 1)]


def repair(ws, lint, cfg):
    """Apply the repair for one lint; returns a description of the cells changed."""
    t = lint["type"]
    if t == "range_short_of_table":
        short = lint["range"]
        names = lint.get("named_ranges_covering_block") or []
        full = names[0] if names else lint["table_block"]
        changed = 0
        for cell in _cells(ws, lint["column"], cfg["first_data_row"], cfg["last_data_row"]):
            v = cell.value
            if isinstance(v, str) and short in v:
                cell.value = v.replace(short, full)
                changed += 1
        if not changed:
            raise SystemExit("A1-type repair found no formula using %s" % short)
        return {"cells": "%s%d:%s%d" % (lint["column"], cfg["first_data_row"], lint["column"], cfg["last_data_row"]),
                "cells_changed": changed, "repair": "%s -> %s" % (short, full)}
    if t in ("hardcoded_value_in_formula_column", "inconsistent_formula"):
        ref = "%s%d" % (lint["column"], lint["row"])
        before = ws[ref].value
        ws[ref].value = lint["column_rule"]
        return {"cells": ref, "cells_changed": 1, "repair": "%r -> %s" % (before, lint["column_rule"])}
    raise SystemExit("no repair defined for lint type %s" % t)


def patch(seed):
    import openpyxl
    from tools.gen_workbook import normalize_xlsx
    decisions = C.load_decisions()
    if not decisions:
        print("patch_workbook: no decisions in decisions/decisions.jsonl; nothing to patch")
        return None
    cfg = C.config()
    lints = {l["id"]: l for l in C.lints()}
    src = os.path.join(C.ROOT, cfg["workbook"])
    wb = openpyxl.load_workbook(src)
    ws = wb[cfg["sheets"]["calc"]]
    patches = []
    for did in sorted(decisions):
        d = decisions[did]
        lint = C.lint_for_decision_id(did)
        if d.get("lint") and d["lint"] != lint["id"]:
            raise SystemExit("%s names lint %s but maps to %s" % (did, d["lint"], lint["id"]))
        if d.get("option") != "adopt-manual":
            patches.append({"decision": did, "lint": lint["id"], "option": d.get("option"), "cells_changed": 0})
            continue
        want = REPAIR_RULES.get(lint["type"])
        if d.get("rule") != want or lints[lint["id"]].get("manual_rule") != want:
            raise SystemExit("%s cites %s but the %s repair implements %s" % (did, d.get("rule"), lint["type"], want))
        p = repair(ws, lint, cfg)
        p.update(decision=did, lint=lint["id"], option="adopt-manual", rule=want, type=lint["type"])
        patches.append(p)
    cp = C.cache_paths(seed)
    os.makedirs(os.path.dirname(cp["patched_book"]), exist_ok=True)
    wb.save(cp["patched_book"])
    normalize_xlsx(cp["patched_book"])
    _, policies = C.load_inputs(seed)
    info = E.expand(cp["patched_book"], policies, cp["expanded_patched"])
    man = {"seed": seed, "decisions_sha256": C.sha256_file(C.DECISIONS),
           "workbook_sha256": C.sha256_file(src), "patched_workbook_sha256": C.sha256_file(cp["patched_book"]),
           "expanded_patched_sha256": C.sha256_file(cp["expanded_patched"]), "patches": patches,
           "expand_seconds": info["seconds"]}
    C.write_json(cp["patch_manifest"], man)
    return man


def main(argv=None):
    ap = argparse.ArgumentParser(description="Apply adopt-manual decisions and expand the patched book")
    ap.add_argument("--seed", type=int, default=C.DEFAULT_SEED)
    ap.add_argument("--service", default=C.DEFAULT_SERVICE, help="accepted for uniformity; unused")
    a = ap.parse_args(argv)
    man = patch(a.seed)
    if man:
        print(json.dumps({"patches": man["patches"], "expanded": C.rel(C.cache_paths(a.seed)["expanded_patched"])}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
